"""Generate the electrical safety standards report as a Word document."""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT = Path(__file__).with_name("electrical_safety_standards_report.docx")


def shade_cell(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_bullets(document, items):
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_standard_table(document):
    rows = [
        ("IEEE 1584-2018", "Guide for arc-flash hazard calculations", "Incident energy, arc-flash boundaries, and equipment studies"),
        ("IEC 60364-4-41 / 6", "Shock protection and verification of LV installations", "LV design, protective measures, inspection, and testing"),
        ("IEC 61140", "Common protection principles against electric shock", "Basic protection, fault protection, bonding, and conductors"),
        ("NFPA 70 (NEC)", "US electrical installation code", "Wiring, overcurrent protection, grounding, and installation"),
        ("NFPA 70E", "US workplace electrical safety practice", "Safe work condition, risk assessment, boundaries, and PPE"),
        ("OSHA 29 CFR 1910", "US workplace electrical regulation", "Employer duties, qualified persons, guarding, and practices"),
        ("ISO 45001", "Occupational health and safety management system", "Hazard controls, competence, audits, and improvement"),
        ("EN 50110-1", "Operation of electrical installations", "Safe operation and electrical work in the European context"),
        ("UL 508A / UL 489", "Product and industrial control-panel safety", "Panel construction, components, and circuit protection"),
        ("IEEE C2 (NESC)", "Safety rules for utility lines and facilities", "Generation, transmission, distribution, and communications"),
    ]
    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, ("Standard or regulation", "Primary purpose", "Typical application")):
        set_cell_text(cell, header, bold=True, color=(255, 255, 255))
        shade_cell(cell, "1F4E79")
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            set_cell_text(cell, value)


def add_reference(document, title, url):
    paragraph = document.add_paragraph(style="List Number")
    paragraph.add_run(title + ". ").bold = True
    paragraph.add_run(url)


def build_report():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for style_name, size, color in (("Title", 24, "1F4E79"), ("Heading 1", 16, "1F4E79"), ("Heading 2", 12, "2F5597")):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Electrical Safety Standards in Industry")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("IEEE, IEC, NFPA, OSHA, ISO, EN, UL, and NESC frameworks").italic = True
    date_paragraph = document.add_paragraph()
    date_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_paragraph.add_run(f"Prepared for electrical engineering study | {date.today().isoformat()}")
    document.add_page_break()

    document.add_heading("Executive Summary", level=1)
    document.add_paragraph("Electrical safety standards convert broad safety goals into repeatable design rules, work practices, verification activities, and management controls. No single document covers every hazard. Engineers normally combine installation standards, product standards, workplace regulations, and an occupational safety management system.")
    document.add_paragraph("IEEE standards are especially valuable for engineering analysis and power-system practice; IEC standards provide internationally harmonized installation and protection principles; NFPA 70 and NFPA 70E are widely used in the United States for installation and energized-work controls; OSHA establishes enforceable US workplace duties. ISO 45001 supplies the management-system structure that keeps controls active over the life of an asset.")
    document.add_paragraph("The central conclusion is that compliance is a system, not a label. Safety depends on competent people, correctly selected equipment, protective coordination, grounding and bonding, documented risk assessment, de-energization, verification, maintenance, and learning from incidents and near misses.")

    document.add_heading("1. Purpose and Scope", level=1)
    document.add_paragraph("This report surveys major electrical safety standards relevant to industrial facilities, laboratories, power systems, control panels, and low-voltage equipment. It explains the purpose and typical application of each framework, compares their roles, and proposes an implementation method for a small 230 V AC to 5 V DC power-supply project. It is educational guidance, not a substitute for the legally applicable edition, an authority having jurisdiction, or a qualified electrical safety professional.")

    document.add_heading("2. How the Standards System Fits Together", level=1)
    document.add_paragraph("Electrical safety documents have different legal and technical roles:")
    add_bullets(document, [
        "Regulations: Government rules such as OSHA requirements may be legally enforceable for covered employers and workers.",
        "Codes: Installation codes such as NFPA 70 are adopted or referenced by jurisdictions and define minimum construction and installation requirements.",
        "Consensus standards and guides: IEEE and IEC documents provide methods, terminology, test procedures, and engineering good practice; their legal force depends on adoption, contract, or policy.",
        "Product standards: UL and IEC product standards address the safety of particular equipment types and their construction, testing, and certification.",
        "Management systems: ISO 45001 organizes leadership, worker participation, risk controls, competence, audits, corrective action, and continual improvement.",
    ])
    add_standard_table(document)

    document.add_heading("3. Major Standards and Their Applications", level=1)
    document.add_heading("3.1 IEEE standards", level=2)
    document.add_paragraph("IEEE develops technical standards and guides used by power, electronics, and controls professionals. IEEE 1584-2018 provides a calculation method for incident energy and arc-flash boundaries based on specified equipment and system parameters. It supports an arc-flash study; it does not by itself select all PPE or replace a workplace safety program. IEEE guides such as IEEE 3007.3 support industrial and commercial power-system safety practice, while IEEE C2, the National Electrical Safety Code, addresses utility supply and communication facilities.")
    add_bullets(document, [
        "Application: collect short-circuit current, clearing time, enclosure, electrode configuration, conductor gap, and working distance data; calculate incident energy and boundaries; label equipment and control exposure.",
        "Engineering value: provides a documented, repeatable analysis that connects protection-system performance to worker exposure.",
        "Limitation: results are only as reliable as the input data, protective-device settings, maintenance condition, and selected edition or jurisdictional rule set.",
    ])

    document.add_heading("3.2 IEC standards", level=2)
    document.add_paragraph("IEC 60364-4-41 sets protection principles against electric shock for low-voltage installations, including basic protection and fault protection. IEC 60364-6 addresses initial and periodic verification. IEC 61140 provides common principles for protection against electric shock across equipment and installations, including protective bonding and automatic disconnection concepts. IEC 60479-1 describes the effects of current on people and livestock, supporting shock-risk assessment. IEC 61482-2 addresses protective clothing against the thermal hazards of electric arc.")
    add_bullets(document, [
        "Application: design insulation, barriers, enclosures, protective conductors, equipotential bonding, automatic disconnection, and verification tests.",
        "Engineering value: establishes internationally consistent vocabulary and protection concepts that can be applied across manufacturers and countries.",
        "Limitation: national committees may publish modifications or complementary rules; the applicable national adoption must be checked.",
    ])

    document.add_heading("3.3 NFPA 70 and NFPA 70E", level=2)
    document.add_paragraph("NFPA 70, the National Electrical Code, governs electrical installation safety where adopted by a jurisdiction or required by contract. It covers wiring methods, overcurrent protection, grounding and bonding, equipment installation, and special occupancies. NFPA 70E focuses on employee safety during electrical work: risk assessment, establishing an electrically safe work condition, shock and arc-flash approach boundaries, job planning, training, and PPE.")
    document.add_paragraph("The distinction matters: NFPA 70 primarily controls how an installation is built, while NFPA 70E controls how workers interact with electrical hazards. A compliant installation can still be dangerous if energized work is poorly planned or equipment is not maintained.")

    document.add_heading("3.4 OSHA electrical safety rules", level=2)
    document.add_paragraph("In the United States, OSHA 29 CFR Part 1910 Subpart S contains general industry electrical requirements, including installation and use of electrical equipment. Sections 1910.331 through 1910.335 address safety-related work practices, qualified and unqualified persons, training, selection and use of work practices, and personal protective equipment. Section 1910.333 emphasizes de-energizing exposed live parts before employees work on or near them, unless specific exceptions apply.")
    document.add_paragraph("OSHA is a regulatory floor, not a complete design handbook. Employers must assess hazards, train workers, maintain equipment, control access, and document a program appropriate to the work. Consensus standards may be used as evidence of recognized good practice, but the employer must determine which requirements are legally applicable.")

    document.add_heading("3.5 ISO 45001", level=2)
    document.add_paragraph("ISO 45001 specifies requirements for an occupational health and safety management system. It is not an electrical installation code. Its value is organizational: leadership and worker participation, hazard identification, risk and opportunity assessment, operational planning, competence, emergency preparedness, performance evaluation, incident investigation, and continual improvement.")
    document.add_paragraph("For electrical safety, ISO 45001 can govern how an organization identifies electrical hazards, assigns owners, verifies controls, manages contractors and changes, and audits recurring tasks such as switching, testing, and maintenance.")

    document.add_heading("3.6 EN 50110-1, UL standards, and NESC", level=2)
    document.add_paragraph("EN 50110-1 gives principles for the safe operation of electrical installations and electrical work, including responsibilities, work organization, isolation, and procedures. UL 508A is widely used for industrial control panels; UL 489 addresses molded-case circuit breakers and circuit-breaker enclosures. Their use supports safe equipment construction and evaluated components, but certification scope and the actual listing conditions must be respected. IEEE C2/NESC provides utility-sector rules for safeguarding workers and the public near electric supply and communication lines.")

    document.add_heading("4. Why These Standards Matter", level=1)
    add_bullets(document, [
        "Prevent fatal shock and burns: insulation, barriers, bonding, automatic disconnection, approach boundaries, and PPE reduce exposure to current and arc energy.",
        "Control fire and equipment damage: correct conductor sizing, overcurrent protection, short-circuit ratings, and installation methods limit overheating and fault escalation.",
        "Make safety measurable: inspection, testing, labels, settings records, and audits turn intent into evidence.",
        "Create a shared language: common definitions for qualified persons, safe work conditions, protective devices, and risk controls reduce ambiguity between designers, operators, contractors, and inspectors.",
        "Support resilient operations: maintenance and change control keep protective settings, enclosures, interlocks, and emergency procedures aligned with the actual system.",
        "Enable legal and professional accountability: following adopted codes and recognized engineering methods helps demonstrate due diligence, while never replacing competent judgment.",
    ])

    document.add_heading("5. Recommended Implementation Workflow", level=1)
    steps = [
        ("Define jurisdiction and scope", "Identify the country, authority having jurisdiction, facility type, voltage level, equipment boundaries, adopted code edition, and contractual requirements."),
        ("Identify hazards", "Consider shock, arc flash, arc blast, fire, stored energy, backfeed, capacitors, rotating equipment, battery systems, electromagnetic effects, and unexpected energization."),
        ("Select the controlling documents", "Map installation, product, workplace, utility, and management-system requirements. Record the selected edition and resolve conflicts through the responsible engineer or authority."),
        ("Design out risk", "Use de-energized operation, insulation, barriers, interlocks, grounding and bonding, current-limiting protection, suitable interrupt ratings, and physical separation before relying on PPE."),
        ("Verify and commission", "Perform continuity, insulation resistance where appropriate, polarity, protective-device, grounding, functional, and trip tests using qualified personnel and calibrated instruments."),
        ("Plan work", "Use job briefings, one-line diagrams, switching instructions, lockout/tagout, test-before-touch, boundaries, permits, PPE, and rescue/emergency arrangements."),
        ("Maintain and improve", "Inspect, clean, exercise, and test equipment; update labels and studies after changes; investigate incidents and near misses; audit the program and retrain as needed."),
    ]
    for heading, text in steps:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(f"{heading}: ").bold = True
        paragraph.add_run(text)

    document.add_heading("6. Application to the Project Power Supply", level=1)
    document.add_paragraph("For the project circuit (230 V AC supply, transformer, bridge rectifier, capacitor, 7805 regulator, and 5 V load), the following controls should be documented before construction or testing:")
    add_bullets(document, [
        "Enclose the 230 V primary and provide a correctly rated disconnect, fuse or breaker, strain relief, insulation, and protective earthing where required.",
        "Keep the low-voltage secondary physically separated from the primary; identify SELV or other circuit classification only after verifying the applicable IEC or national requirements.",
        "Use a transformer and components with suitable voltage, current, temperature, insulation, and short-circuit ratings; do not infer safety from the 5 V output alone.",
        "Provide discharge control or a documented wait-and-measure procedure for the filter capacitor; capacitors can remain hazardous after power removal.",
        "Use barriers and a covered enclosure during energized measurements. Only qualified persons should access the primary side, with an appropriate risk assessment and test equipment.",
        "Verify protective earth continuity, polarity, insulation, absence of exposed live parts, fuse rating, and output voltage before connecting a load.",
        "Keep a schematic, bill of materials, test results, hazard assessment, and revision record. Reassess safety after component, enclosure, transformer, or protection changes.",
    ])

    document.add_heading("7. Limitations and Good Practice", level=1)
    document.add_paragraph("Standards change, are copyrighted, and are often adopted with amendments. Titles and examples in this report are orientation points rather than permission to reproduce or substitute for the normative text. The applicable authority, current edition, local law, equipment listing, and manufacturer instructions must be checked before design, installation, or work. A risk assessment should account for the real installation, not merely the nominal voltage.")

    document.add_heading("8. References and Further Research", level=1)
    refs = [
        ("OSHA, 29 CFR 1910.333, Selection and use of work practices", "https://www.osha.gov/laws-regs/regulations/standardnumber/1910/1910.333"),
        ("OSHA, 29 CFR 1910 Subpart S, Electrical", "https://www.osha.gov/laws-regs/regulations/standardnumber/1910/1910subpartS"),
        ("NFPA, NFPA 70: National Electrical Code", "https://www.nfpa.org/codes-and-standards/nfpa-70-standard-development/70"),
        ("NFPA, NFPA 70E: Standard for Electrical Safety in the Workplace", "https://www.nfpa.org/codes-and-standards/nfpa-70e-standard-development/70e"),
        ("IEEE Standards Association, IEEE 1584-2018", "https://standards.ieee.org/standard/1584-2018.html"),
        ("IEEE Standards Association, National Electrical Safety Code (IEEE C2)", "https://standards.ieee.org/products-programs/nesc/"),
        ("IEC, IEC 60364 series: Low-voltage electrical installations", "https://webstore.iec.ch/en/publication/6029"),
        ("IEC, IEC 61140: Protection against electric shock", "https://webstore.iec.ch/en/publication/6437"),
        ("ISO, ISO 45001: Occupational health and safety management systems", "https://www.iso.org/iso-45001-occupational-health-and-safety.html"),
        ("UL Standards and Engagement, UL 508A", "https://www.shopulstandards.com/ProductDetail.aspx?productId=UL508A"),
    ]
    for title_text, url in refs:
        add_reference(document, title_text, url)
    document.add_paragraph(f"Sources accessed: {date.today().isoformat()}. Always confirm the current edition and local adoption before relying on a requirement.")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Electrical Safety Standards Report | Educational use")
    document.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build_report()